"""
QualityMetrics月次実績 自動更新（Python版・コア処理）
データ提供者の生データ(Monthlyレポート) → QualityMetrics実績.xlsx working file への転記

設計方針：
- 列位置はハードコードせずヘッダー文字列検索で特定（構造変更に強くする）
- 順方向チェック：working file側の行が元データに一致するか（＝③店舗マスタ変更の反映漏れ検知）
- 逆方向チェック：元データにあってworking file側に行がない＝新規コード追加が必要な可能性
  （今回、徳島の新コード86010で実際にこのケースを確認したため追加した）
"""
import argparse
import sys
import openpyxl
from pathlib import Path
from datetime import date


def default_target_month(today=None):
    """対象月＝当月ではなく「先月」がデフォルト。
    月初に前月分のレポートを処理する運用のため（0701_QualityMetrics.txt 00:04:50〜と整合）。
    """
    today = today or date.today()
    return 12 if today.month == 1 else today.month - 1


def _check_source_file(path):
    """データ提供者のMonthlyレポートとして妥当か検証する。(valid, reason)を返す（例外を投げない）。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        sheets = set(wb.sheetnames)
        wb.close()
    except Exception as e:
        return False, f"開けませんでした（{e}）"

    if "今月Qチェック状況" in sheets:
        return True, "OK"
    if {"店舗毎", "AM毎"} & sheets:
        return False, "別の週次データ（IMM＝是正対応トラッキング）です。Monthly_Reportではありません"
    return False, f"「今月Qチェック状況」シートがありません（実際: {', '.join(sorted(sheets))}）"


def _check_working_file(path, target_sheet="QualityMetrics実績"):
    """working fileとして妥当か検証する。(valid, reason)を返す（例外を投げない）。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        sheets = set(wb.sheetnames)
        wb.close()
    except Exception as e:
        return False, f"開けませんでした（{e}）"

    if target_sheet in sheets:
        return True, "OK"
    return False, f"「{target_sheet}」シートがありません（実際: {', '.join(sorted(sheets))}）"


def resolve_input(path_str, checker_fn, kind_label):
    """
    path_strがファイルならそのまま検証して返す。
    フォルダなら中の.xlsxを"全件"スキャンし、妥当なものを自動選択する。
    ・妥当なものが1件だけ → それを自動採用（複数ファイルが混在していても正しい1件を拾う）
    ・0件 → 見つかった全ファイルとそれぞれの却下理由を提示してエラー
    ・2件以上 → 全候補を提示したうえで、最終更新日時が最も新しいものを暫定採用（要確認の警告付き）
    """
    p = Path(path_str)

    if p.is_file():
        ok, reason = checker_fn(str(p))
        if not ok:
            raise FileNotFoundError(f"「{p.name}」は{kind_label}として使えません：{reason}")
        return str(p)

    if not p.is_dir():
        raise FileNotFoundError(f"「{path_str}」が見つかりません。")

    # Excelのロックファイル(~$から始まる一時ファイル)は除外
    candidates = sorted(c for c in p.glob("*.xlsx") if not c.name.startswith("~$"))

    if not candidates:
        raise FileNotFoundError(f"「{p}」フォルダにxlsxファイルがありません。{kind_label}を入れてください。")

    valid, rejected = [], []
    for c in candidates:
        ok, reason = checker_fn(str(c))
        (valid if ok else rejected).append((c, reason))

    if len(valid) == 1:
        chosen = valid[0][0]
        if len(candidates) > 1:
            print(f"[自動選択] {kind_label}: 「{chosen.name}」を採用（フォルダ内の他{len(candidates)-1}件は対象外と判定）")
        return str(chosen)

    if len(valid) == 0:
        msg = f"「{p}」フォルダ内に{kind_label}として使えるファイルがありません。\n"
        msg += f"フォルダ内の{len(candidates)}件を全て確認しましたが、いずれも該当しませんでした：\n"
        msg += "\n".join(f"  - {c.name}：{reason}" for c, reason in rejected)
        raise FileNotFoundError(msg)

    # 妥当な候補が複数：最終更新日時が新しい順に並べ、最新を暫定採用
    valid_sorted = sorted(valid, key=lambda x: x[0].stat().st_mtime, reverse=True)
    chosen, others = valid_sorted[0][0], valid_sorted[1:]
    print(f"[警告] {kind_label}として使えるファイルが{len(valid)}件見つかりました。"
          f"最終更新日時が最も新しい「{chosen.name}」を暫定採用します。")
    print("  他の候補: " + ", ".join(c.name for c, _ in others))
    print(f"  意図と違う場合は、使わないファイルを「{p}」から出してから再実行してください。")
    return str(chosen)


def _find_header_cell(ws, text, max_row=8, max_col=100):
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if v and text in str(v):
                return r, c
    return None, None


def load_qf_dict(src_path, sheet_name="今月Qチェック状況"):
    wb = openpyxl.load_workbook(src_path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    hr, code_col = _find_header_cell(ws, "ワーク", max_row=5)
    _, total_col = _find_header_cell(ws, "合計値", max_row=5)
    if hr is None or total_col is None:
        raise ValueError(f"{sheet_name}: ヘッダーが想定と異なります")
    q_col, f_col = total_col, total_col + 1
    dictQ, dictF = {}, {}
    for row in ws.iter_rows(min_row=hr + 2, values_only=False):
        code = row[code_col - 1].value
        if code:
            dictQ[str(code).strip()] = row[q_col - 1].value
            dictF[str(code).strip()] = row[f_col - 1].value
    wb.close()
    return dictQ, dictF


def load_mycheck_dict(src_path, sheet_name="マイチェック状況"):
    wb = openpyxl.load_workbook(src_path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    hr, code_col = _find_header_cell(ws, "ワーク", max_row=6)
    _, date_col = _find_header_cell(ws, "直近実施日", max_row=6)
    if hr is None or date_col is None:
        raise ValueError(f"{sheet_name}: ヘッダーが想定と異なります")
    dictMy = {}
    for row in ws.iter_rows(min_row=hr + 1, values_only=False):
        code = row[code_col - 1].value
        if code:
            dictMy[str(code).strip()] = row[date_col - 1].value
    wb.close()
    return dictMy


def update_working_file(working_path, dictQ, dictF, dictMy, target_month, output_path,
                         target_sheet="QualityMetrics実績", code_header="DC", finalcheck_header="最終チェック",
                         mydate_header="直近実施日"):
    wb = openpyxl.load_workbook(working_path)  # 書き込みのためread_only=False
    ws = wb[target_sheet]

    hr, code_col = _find_header_cell(ws, code_header, max_row=6)
    _, final_col = _find_header_cell(ws, finalcheck_header, max_row=6)
    if hr is None or final_col is None:
        raise ValueError(f"{target_sheet}: ヘッダーが想定と異なります")

    month_label = f"{target_month}月"

    def find_month_col(col_start, col_end, label):
        for c in range(col_start, col_end + 1):
            v = ws.cell(row=hr, column=c).value
            if v and label == str(v).strip():
                return c
        return None

    q_month_col = find_month_col(code_col, final_col - 1, month_label)
    f_month_col = find_month_col(final_col, ws.max_column, month_label)
    my_col = None
    for c in range(final_col, ws.max_column + 1):
        v = ws.cell(row=hr, column=c).value
        if v and mydate_header in str(v):
            my_col = c
            break

    if not (q_month_col and f_month_col and my_col):
        raise ValueError(f"当月列（{month_label}）またはマイチェック列が見つかりません")

    updated, mismatches = 0, []
    consumed_codes = set()

    for r in range(hr + 1, ws.max_row + 1):
        code = ws.cell(row=r, column=code_col).value
        if code in (None, ""):
            continue
        code = str(code).strip()
        if code in dictQ:
            ws.cell(row=r, column=q_month_col).value = dictQ[code]
            ws.cell(row=r, column=f_month_col).value = dictF[code]
            updated += 1
            consumed_codes.add(code)
        else:
            store_name = ws.cell(row=r, column=code_col + 1).value
            mismatches.append((code, store_name))
        if code in dictMy:
            ws.cell(row=r, column=my_col).value = dictMy[code]

    # 逆方向チェック：元データにあってworking fileに行がないコード（新規店舗の可能性）
    new_candidates = set(dictQ.keys()) - consumed_codes

    wb.save(output_path)
    wb.close()
    return {
        "updated": updated,
        "mismatches": mismatches,
        "new_store_candidates": sorted(new_candidates),
    }


def main():
    p = argparse.ArgumentParser(description="QualityMetrics月次実績：データ提供者データをworking fileへ転記")
    p.add_argument("--source", required=True,
                    help="データ提供者のMonthlyレポート(xlsx)、またはそれが入ったフォルダのパス")
    p.add_argument("--working", required=True,
                    help="QualityMetrics実績.xlsx（working file）、またはそれが入ったフォルダのパス")
    p.add_argument("--output", required=True, help="出力先パス（working fileを直接上書きしない）")
    p.add_argument("--month", type=int, default=None,
                    help="対象月（1-12）。省略時は「先月」を自動採用")
    args = p.parse_args()

    try:
        source_path = resolve_input(args.source, _check_source_file, "データ提供者のMonthlyレポート")
        working_path = resolve_input(args.working, _check_working_file, "working file")

        target_month = args.month or default_target_month()
        print(f"対象月: {target_month}月")
        print(f"データ提供者レポート: {source_path}")
        print(f"working file  : {working_path}")

        dictQ, dictF = load_qf_dict(source_path)
        dictMy = load_mycheck_dict(source_path)
        print(f"データ提供者データ件数: Q={len(dictQ)}件 マイチェック={len(dictMy)}件")

        result = update_working_file(working_path, dictQ, dictF, dictMy,
                                      target_month=target_month, output_path=args.output)
        print(f"転記件数: {result['updated']}")
        if result["mismatches"]:
            print(f"【要確認】working fileにあるがデータ提供者データにないコード（③店舗マスタ変更の反映漏れの可能性）:")
            for code, name in result["mismatches"]:
                print(f"  - {code} {name}")
        if result["new_store_candidates"]:
            print(f"【要確認】データ提供者データにあるがworking fileに行がないコード（新規店舗追加の可能性）:")
            for code in result["new_store_candidates"]:
                print(f"  - {code}")
        if not result["mismatches"] and not result["new_store_candidates"]:
            print("突合せ：完全一致（要確認なし）")
        print(f"出力: {args.output}")
        return 0

    except FileNotFoundError as e:
        print(f"\n【入力エラー】{e}\n")
        return 1
    except Exception as e:
        print(f"\n【エラー】想定外の問題が発生しました: {e}")
        print("このメッセージをそのままの形でご相談ください。")
        return 1


if __name__ == "__main__":
    sys.exit(main() or 0)
