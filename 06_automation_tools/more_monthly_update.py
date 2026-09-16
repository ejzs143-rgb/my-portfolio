"""
MORE Partner分析 月次更新スクリプト
=====================================
使い方:
  python more_monthly_update.py --vw VW_Partner.xlsx --audi Audi_Partner.xlsx --month 2026/01

初回: 新規Excelを生成
2回目以降: 既存Excelの推移シートに新月列を追加 + ワースト店シート更新
"""

import argparse, json, os, sys
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# === Styles ===
HEADER_FILL = PatternFill('solid', fgColor='2F5496')
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
AVG_FILL = PatternFill('solid', fgColor='FFF2CC')
AVG_FONT = Font(name='Arial', bold=True, size=9)
DATA_FONT = Font(name='Arial', size=9)
THIN_BORDER = Border(
    left=Side(style='thin', color='B0B0B0'), right=Side(style='thin', color='B0B0B0'),
    top=Side(style='thin', color='B0B0B0'), bottom=Side(style='thin', color='B0B0B0'))
SCORE_HIGH = PatternFill('solid', fgColor='F4CCCC')
SCORE_MED = PatternFill('solid', fgColor='FCE5CD')
SCORE_LOW = PatternFill('solid', fgColor='D9EAD3')
PCT_FMT = '0.0'
SCORE_FMT = '0.00'


def sc(ws, row, col, value, font=DATA_FONT, fill=None, fmt=None, align=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    cell.border = THIN_BORDER
    if fill: cell.fill = fill
    if fmt: cell.number_format = fmt
    if align: cell.alignment = align
    return cell


def extract(filepath, brand):
    """Partner Excelから3指標+全ディーラーデータを抽出"""
    df = pd.read_excel(filepath, sheet_name='Process', header=None)

    # P指標の列位置を自動検出（row 8）
    row8 = df.iloc[8]
    p_cols = {}
    for i, v in enumerate(row8):
        if pd.notna(v) and str(v).strip().startswith('P0'):
            p_cols[str(v).strip()] = i

    avgs = {}
    for p, col in p_cols.items():
        val = df.iloc[11, col]
        if pd.notna(val):
            avgs[p] = float(val)

    dealers = []
    for i in range(15, df.shape[0]):
        pid = df.iloc[i, 1]
        if pd.isna(pid):
            continue
        name = str(df.iloc[i, 2]) if pd.notna(df.iloc[i, 2]) else str(pid)
        d = {
            'partner_id': str(int(pid)) if isinstance(pid, float) else str(pid),
            'name': name,
            'class': str(df.iloc[i, 0]) if pd.notna(df.iloc[i, 0]) else '',
        }
        for p in ['P004', 'P002', 'P016', 'P001', 'P019', 'P005']:
            if p in p_cols:
                val = df.iloc[i, p_cols[p]]
                d[p] = float(val) if pd.notna(val) else 0.0
            else:
                d[p] = 0.0

        score = 0
        for p in ['P004', 'P002', 'P016']:
            if p in avgs and avgs[p] > 0:
                score += d[p] / avgs[p]
        d['score'] = round(score, 2)
        dealers.append(d)

    dealers.sort(key=lambda x: -x['score'])
    print(f"[{brand}] {len(dealers)} dealers extracted. Top: {dealers[0]['name']} (Score={dealers[0]['score']})")
    return {'brand': brand, 'avgs': avgs, 'dealers': dealers}


def write_worst_sheet(ws, data, brand, month):
    """ワースト店シートを(再)作成"""
    # Clear existing content
    for merge in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merge))
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None

    avgs = data['avgs']
    dealers = data['dealers']

    ws.merge_cells('A1:L1')
    c = ws.cell(row=1, column=1, value=f'{brand} MORE Partner分析 — {month}（過去12ヶ月ローリング）')
    c.font = Font(name='Arial', bold=True, size=13, color='2F5496')

    ws.merge_cells('A3:L3')
    ws.cell(row=3, column=1,
            value='スコア = P004平均比 + P002平均比 + P016平均比（各指標÷全国平均の合算）').font = \
        Font(name='Arial', size=8, italic=True, color='666666')

    headers = [
        ('順位', 5), ('Partner ID', 9), ('ディーラー名', 35), ('Class', 5),
        ('P004\nKEP比率%', 10), ('平均比', 7), ('P002\nGW件数%', 10), ('平均比', 7),
        ('P016\nLast6m%', 10), ('平均比', 7), ('スコア', 8), ('判定', 10),
        ('P001\nGWコスト%', 10), ('P019\nKEPコスト%', 10), ('P005\nAタイム%', 10),
    ]
    for i, (h, w) in enumerate(headers):
        sc(ws, 4, i + 1, h, font=HEADER_FONT, fill=HEADER_FILL,
           align=Alignment(horizontal='center', vertical='center', wrap_text=True))
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.row_dimensions[4].height = 35

    # Average row
    sc(ws, 5, 3, '全国平均', fill=AVG_FILL, font=AVG_FONT)
    for col in [1, 2, 4]:
        sc(ws, 5, col, '', fill=AVG_FILL, font=AVG_FONT)
    sc(ws, 5, 5, avgs.get('P004', 0), fill=AVG_FILL, font=AVG_FONT, fmt=PCT_FMT)
    sc(ws, 5, 6, 1.00, fill=AVG_FILL, font=AVG_FONT, fmt=SCORE_FMT)
    sc(ws, 5, 7, avgs.get('P002', 0), fill=AVG_FILL, font=AVG_FONT, fmt=PCT_FMT)
    sc(ws, 5, 8, 1.00, fill=AVG_FILL, font=AVG_FONT, fmt=SCORE_FMT)
    sc(ws, 5, 9, avgs.get('P016', 0), fill=AVG_FILL, font=AVG_FONT, fmt=PCT_FMT)
    sc(ws, 5, 10, 1.00, fill=AVG_FILL, font=AVG_FONT, fmt=SCORE_FMT)
    sc(ws, 5, 11, 3.00, fill=AVG_FILL, font=AVG_FONT, fmt=SCORE_FMT)
    sc(ws, 5, 12, '（基準）', fill=AVG_FILL, font=AVG_FONT)
    sc(ws, 5, 13, avgs.get('P001', 0), fill=AVG_FILL, font=AVG_FONT, fmt=PCT_FMT)
    sc(ws, 5, 14, avgs.get('P019', 0), fill=AVG_FILL, font=AVG_FONT, fmt=PCT_FMT)
    sc(ws, 5, 15, avgs.get('P005', 0), fill=AVG_FILL, font=AVG_FONT, fmt=PCT_FMT)

    for idx, d in enumerate(dealers):
        r = 6 + idx
        sc(ws, r, 1, idx + 1, align=Alignment(horizontal='center'))
        sc(ws, r, 2, d['partner_id'], font=Font(name='Arial', size=8))
        sc(ws, r, 3, d['name'], font=Font(name='Arial', size=8))
        sc(ws, r, 4, d.get('class', ''), align=Alignment(horizontal='center'))
        sc(ws, r, 5, d['P004'], fmt=PCT_FMT)
        sc(ws, r, 6, f'=IF(E$5=0,0,E{r}/E$5)', fmt=SCORE_FMT)
        sc(ws, r, 7, d['P002'], fmt=PCT_FMT)
        sc(ws, r, 8, f'=IF(G$5=0,0,G{r}/G$5)', fmt=SCORE_FMT)
        sc(ws, r, 9, d['P016'], fmt=PCT_FMT)
        sc(ws, r, 10, f'=IF(I$5=0,0,I{r}/I$5)', fmt=SCORE_FMT)
        sc(ws, r, 11, f'=F{r}+H{r}+J{r}', fmt=SCORE_FMT)
        sc(ws, r, 12, f'=IF(K{r}>=8,"★★★要指導",IF(K{r}>=5,"★★要注意",IF(K{r}>=3,"★注目","")))')
        sc(ws, r, 13, d.get('P001', 0), fmt=PCT_FMT)
        sc(ws, r, 14, d.get('P019', 0), fmt=PCT_FMT)
        sc(ws, r, 15, d.get('P005', 0), fmt=PCT_FMT)
        if d['score'] >= 8:
            ws.cell(row=r, column=11).fill = SCORE_HIGH
            ws.cell(row=r, column=12).fill = SCORE_HIGH
        elif d['score'] >= 5:
            ws.cell(row=r, column=11).fill = SCORE_MED
            ws.cell(row=r, column=12).fill = SCORE_MED
        elif d['score'] < 3:
            ws.cell(row=r, column=11).fill = SCORE_LOW

    ws.freeze_panes = 'A6'
    ws.auto_filter.ref = f'A4:O{5 + len(dealers)}'


def update_transition_sheet(ws, vw_data, audi_data, month):
    """推移シートに新月列を追加"""
    vw_header_row = 4
    next_col = 4
    while ws.cell(row=vw_header_row, column=next_col).value is not None:
        if ws.cell(row=vw_header_row, column=next_col).value == month:
            break
        next_col += 1
    ws.column_dimensions[get_column_letter(next_col)].width = 10
    sc(ws, vw_header_row, next_col, month, font=HEADER_FONT, fill=HEADER_FILL,
       align=Alignment(horizontal='center'))
    is_first = (next_col == 4)
    for idx, d in enumerate(vw_data['dealers'][:20]):
        r = 5 + idx
        if is_first:
            sc(ws, r, 1, idx + 1, align=Alignment(horizontal='center'))
            sc(ws, r, 2, d['partner_id'], font=Font(name='Arial', size=8))
            sc(ws, r, 3, d['name'], font=Font(name='Arial', size=8))
        sc(ws, r, next_col, d['score'], fmt=SCORE_FMT)
        if d['score'] >= 8: ws.cell(row=r, column=next_col).fill = SCORE_HIGH
        elif d['score'] >= 5: ws.cell(row=r, column=next_col).fill = SCORE_MED
    audi_start = 27
    sc(ws, audi_start + 1, next_col, month, font=HEADER_FONT, fill=HEADER_FILL,
       align=Alignment(horizontal='center'))
    for idx, d in enumerate(audi_data['dealers'][:20]):
        r = audi_start + 2 + idx
        if is_first:
            sc(ws, r, 1, idx + 1, align=Alignment(horizontal='center'))
            sc(ws, r, 2, d['partner_id'], font=Font(name='Arial', size=8))
            sc(ws, r, 3, d['name'], font=Font(name='Arial', size=8))
        sc(ws, r, next_col, d['score'], fmt=SCORE_FMT)
        if d['score'] >= 8: ws.cell(row=r, column=next_col).fill = SCORE_HIGH
        elif d['score'] >= 5: ws.cell(row=r, column=next_col).fill = SCORE_MED


def write_guide_sheet(ws):
    ws.merge_cells('A1:F1')
    ws.cell(row=1, column=1, value='指標解説 & 運用ガイド').font = Font(name='Arial', bold=True, size=13, color='2F5496')
    ws.cell(row=3, column=1, value='■ スコア対象指標（3つ）').font = Font(name='Arial', bold=True, size=10, color='2F5496')
    headers = ['指標', '名称', '意味', '高い場合の示唆', '区分']
    for i, h in enumerate(headers):
        sc(ws, 4, i + 1, h, font=HEADER_FONT, fill=HEADER_FILL)
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 45
    ws.column_dimensions['E'].width = 12
    indicators = [
        ['P004', 'KEP比率（%）', 'GW全体に占めるKEP(CFG)の割合', 'CFG枠の乱用。高い＝なんでもCFGにしてる疑い', '★スコア対象'],
        ['P002', 'GW件数比率（%）', '保証+GW全体に占めるGW件数の割合', 'GWの多用。高い＝保証外をGWで通しすぎ', '★スコア対象'],
        ['P016', 'Last 6m比率（%）', '保証終了前6ヶ月内のクレーム割合', '駆け込み修理。高い＝予防整備で請求してる疑い', '★スコア対象'],
    ]
    for idx, ind in enumerate(indicators):
        for i, v in enumerate(ind):
            fill = PatternFill('solid', fgColor='E2EFDA') if i == 4 else None
            sc(ws, 5 + idx, i + 1, v, fill=fill)
    ws.cell(row=9, column=1, value='■ 参考表示指標（スコア対象外）').font = Font(name='Arial', bold=True, size=10, color='666666')
    ref = [
        ['P001', 'GWコスト比率（%）', 'GW金額の保証+GW全体に対する割合', '金額ベースでのGW比率。P002の補完', '参考'],
        ['P019', 'KEPコスト比率（%）', 'KEP金額のGW全体に対する割合', '金額ベースでのCFG比率。P004の補完', '参考'],
        ['P005', 'Aタイム比率（%）', '全TUに占めるAタイムの割合', 'レイバー効率の目安', '参考'],
    ]
    SUBHEADER_FILL = PatternFill('solid', fgColor='D6E4F0')
    SUBHEADER_FONT = Font(name='Arial', bold=True, size=9)
    for i, h in enumerate(headers):
        sc(ws, 10, i + 1, h, font=SUBHEADER_FONT, fill=SUBHEADER_FILL)
    for idx, ind in enumerate(ref):
        for i, v in enumerate(ind):
            sc(ws, 11 + idx, i + 1, v)
    ws.cell(row=16, column=1, value='■ スコア計算方法').font = Font(name='Arial', bold=True, size=10, color='2F5496')
    for i, t in enumerate([
        'スコア ＝ (P004÷全国平均P004) + (P002÷全国平均P002) + (P016÷全国平均P016)',
        '→ 全指標が平均なら 1.0+1.0+1.0 = 3.0',
        '→ 8以上 ★★★要指導 / 5以上 ★★要注意 / 3以上 ★注目',
        '', '■ データ期間', 'Partner Excelは過去12ヶ月のローリング（例：12月レポート → 2025年1月〜12月）',
        '', '■ 月次運用',
        '1. BPからPartner Excel（VW・Audi）をDL',
        '2. python more_monthly_update.py --vw VW.xlsx --audi Audi.xlsx --month YYYY/MM',
        '3. ワースト店シートが最新月で上書き、推移シートに新月列が追加される',
    ]):
        ws.cell(row=17 + i, column=1, value=t).font = Font(name='Arial', size=9)


def main():
    parser = argparse.ArgumentParser(description='MORE Partner月次分析')
    parser.add_argument('--vw', required=True, help='VW Partner Excel path')
    parser.add_argument('--audi', required=True, help='Audi Partner Excel path')
    parser.add_argument('--month', required=True, help='対象月 (例: 2026/01)')
    parser.add_argument('--output', default='MORE_analysis_system.xlsx', help='出力Excel path')
    args = parser.parse_args()

    print(f"=== MORE月次分析: {args.month} ===")

    vw = extract(args.vw, 'VW')
    audi = extract(args.audi, 'Audi')

    # Save monthly JSON for backup
    json_dir = 'monthly_data'
    os.makedirs(json_dir, exist_ok=True)
    month_key = args.month.replace('/', '')
    with open(f'{json_dir}/vw_{month_key}.json', 'w') as f:
        json.dump(vw, f, ensure_ascii=False)
    with open(f'{json_dir}/audi_{month_key}.json', 'w') as f:
        json.dump(audi, f, ensure_ascii=False)

    # Load or create workbook
    if os.path.exists(args.output):
        print(f"既存ファイルを更新: {args.output}")
        wb = load_workbook(args.output)
    else:
        print(f"新規作成: {args.output}")
        wb = Workbook()
        wb.active.title = 'VW ワースト店'
        wb.create_sheet('Audi ワースト店')
        wb.create_sheet('月次推移')
        wb.create_sheet('指標解説')

    # Update worst sheets (latest month)
    write_worst_sheet(wb['VW ワースト店'], vw, 'VW', args.month)
    write_worst_sheet(wb['Audi ワースト店'], audi, 'Audi', args.month)

    # Update transition sheet
    # Initialize headers if new
    ws_t = wb['月次推移']
    if ws_t.cell(row=1, column=1).value is None:
        ws_t.merge_cells('A1:H1')
        ws_t.cell(row=1, column=1, value='月次スコア推移（ワースト20）').font = \
            Font(name='Arial', bold=True, size=13, color='2F5496')
        ws_t.cell(row=3, column=1, value='VW').font = Font(name='Arial', bold=True, size=11, color='2F5496')
        for i, h in enumerate(['順位', 'Partner ID', 'ディーラー名']):
            sc(ws_t, 4, i + 1, h, font=HEADER_FONT, fill=HEADER_FILL,
               align=Alignment(horizontal='center'))
        ws_t.column_dimensions['A'].width = 5
        ws_t.column_dimensions['B'].width = 9
        ws_t.column_dimensions['C'].width = 35
        ws_t.cell(row=27, column=1, value='Audi').font = Font(name='Arial', bold=True, size=11, color='2F5496')
        for i, h in enumerate(['順位', 'Partner ID', 'ディーラー名']):
            sc(ws_t, 28, i + 1, h, font=HEADER_FONT, fill=HEADER_FILL,
               align=Alignment(horizontal='center'))

    update_transition_sheet(ws_t, vw, audi, args.month)

    # Guide sheet (write once)
    ws_g = wb['指標解説']
    if ws_g.cell(row=1, column=1).value is None:
        write_guide_sheet(ws_g)

    wb.save(args.output)
    print(f"\n完了: {args.output}")
    print(f"  VW: {len(vw['dealers'])}店 / Audi: {len(audi['dealers'])}店")
    print(f"  VW Top3: {', '.join(d['name'].split('-')[-1].strip() for d in vw['dealers'][:3])}")
    print(f"  Audi Top3: {', '.join(d['name'].split('-')[-1].strip() for d in audi['dealers'][:3])}")


if __name__ == '__main__':
    main()
